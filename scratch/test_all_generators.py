import os
import json
import docx
from datetime import datetime
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def test_generators():
    data = {
        "applicant": {
            "fullName": "ROSHAN UPADHYAY",
            "passportNumber": "Y4526175",
            "dateOfBirth": "15/08/1990",
            "placeOfIssue": "MUMBAI",
            "dateOfIssue": "10/01/2020",
            "dateOfExpiry": "09/01/2030",
            "residentialAddress": "Flat 402, Sunshine Heights, Andheri West, Mumbai",
            "city": "Mumbai"
        },
        "travel": {
            "destinationCountry": "France",
            "travelStartDate": "15/10/2026",
            "travelEndDate": "25/10/2026",
            "fundingArrangement": "Self-funded from personal savings",
            "jobTitle": "Senior Software Engineer",
            "employerName": "Infosys Limited",
            "applicantPhone": "+91 9820123456",
            "applicantEmail": "roshan.upadhyay@example.com"
        },
        "travellers": [
            {
                "fullName": "PRIYA UPADHYAY",
                "passportNumber": "Z9876543",
                "relation": "Spouse",
                "occupation": "HR Manager"
            }
        ],
        "hotels": [
            {
                "confirmationNumber": "TBHBV5YP9R",
                "hotelName": "Holiday Inn Paris - Gare de Lyon",
                "leadGuest": "ROSHAN UPADHYAY",
                "checkIn": "15-Oct-2026",
                "checkOut": "25-Oct-2026",
                "duration": "10 Night(s)",
                "city": "Paris",
                "phone": "+33 1 4344 5678",
                "address": "11 Rue de Lyon, 75012 Paris, France",
                "numRooms": "1",
                "numGuests": "2 Adult(s)",
                "guests": [
                    {"guestName": "ROSHAN UPADHYAY", "roomType": "Standard Room", "numGuests": "2 Adult(s)"},
                    {"guestName": "PRIYA UPADHYAY", "roomType": "Standard Room", "numGuests": "2 Adult(s)"}
                ]
            }
        ]
    }

    print("Testing Hotel Blocking Table structure...")
    base_tpl = "templates/hotel-blocking/Single Hotel/Hotel Booking Format.docx"
    doc = docx.Document(base_tpl)
    outer = doc.tables[0]
    sub_outer = outer.rows[0].cells[0].tables[0]
    subtables_cell = sub_outer.rows[1].cells[0]
    sub1 = subtables_cell.tables[1]
    print(f"Original sub1 cols: {len(sub1.columns)}, rows: {len(sub1.rows)}")
    print(f"Header: {[c.text.strip() for c in sub1.rows[0].cells]}")
    print(f"Data: {[c.text.strip() for c in sub1.rows[1].cells]}")

    print("\nTesting Passport Auth Docx...")
    p_auth_tpl = "templates/Passport Authorization Letter/Passport Authorization Letter.docx"
    p_doc = docx.Document(p_auth_tpl)
    print(f"Loaded {p_auth_tpl}, paragraphs: {len(p_doc.paragraphs)}")

    print("\nTesting Company Auth Docx...")
    c_auth_tpl = "templates/Company Authorization Letter/Company Authorization Letter.docx"
    c_doc = docx.Document(c_auth_tpl)
    print(f"Loaded {c_auth_tpl}, paragraphs: {len(c_doc.paragraphs)}")

    print("\nAll templates loaded successfully!")

if __name__ == "__main__":
    test_generators()
